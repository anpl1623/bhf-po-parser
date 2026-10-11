from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from parse_po import (
    MISMATCH_FILL,
    TEMPLATE_PATH,
    build_rows,
    parse_header,
    parse_items,
    parse_totals,
    validate_item,
    validate_totals,
    write_xlsx,
)

SAMPLE = (Path(__file__).parent / "data" / "purchase_order_sample.txt").read_text()


def test_header_fields():
    header = parse_header(SAMPLE)
    assert header["buyer"] == "610 JORDAN MILLER"
    assert header["ship_terms"] == "FOB Ningbo,CHN"
    assert header["po_number"] == "TX9K2QP"
    assert header["ref_master_po"] == "TX9H1LM"
    assert header["ship_date"] == "5/02/26"
    assert header["vendor"] == "50201 HARBOR POINT TRADING LTD, C/O MERIDIAN SOURCING CO INC, 410 COMMERCE AVE, EDISON NJ 08837"
    assert header["ship_to"] == "331 DISTRIBUTION/NORTHEAST, C/O SWIFTLINE EXPRESS, 1450 CROSSROADS BLVD, EDISON, NJ 08817"
    assert header["bill_to"] == "331 MERIDIAN SOURCING CO. INC., C/O NORTHSTAR RETAIL GROUP - DEPT OPS, PO BOX 4021, EDISON, NJ 08837"
    assert header["total_cost"] == Decimal("447455.12")


def test_items_and_upc_pairing():
    items = parse_items(SAMPLE)
    assert len(items) == 12
    first, last = items[0], items[-1]
    assert first["sku"] == "71097188"
    assert first["upc"] == "0091120032129"
    assert first["vendor_part"] == "04212"  # leading zero kept
    assert first["description"] == "NORDPEAK ASHER T CMFST GRY"
    assert first["ext_cost"] == Decimal("6718.400")
    assert last["sku"] == "71293283"
    assert last["upc"] == "0091120032426"
    assert items[7]["description"] == "NORDPEAK ASHER STRIPE F/Q DVT GRAY"


def test_sample_validates_cleanly():
    items = parse_items(SAMPLE)
    assert all(validate_item(item) == {} for item in items)
    assert validate_totals(items, parse_totals(SAMPLE)) == []


# Change the first item's printed EXT QTY from 136 to 137 (same width, so the layout is unchanged).
BROKEN = SAMPLE.replace("   68     2      136", "   68     2      137")


def test_mismatch_is_flagged():
    problems = validate_item(parse_items(BROKEN)[0])
    assert problems == {
        "EXT QTY": "Expected 136 (CTNS x CSPK)",
        "EXT Cost": "Expected 6767.800 (Cost x EXT QTY)",  # 49.400 x 137
    }


def test_mismatch_is_highlighted_in_output(tmp_path):
    items = parse_items(BROKEN)
    rows = build_rows("broken.txt", parse_header(BROKEN), items)
    out_path = tmp_path / "out.xlsx"
    write_xlsx(rows, [validate_item(item) for item in items], out_path)

    sheet = load_workbook(out_path).active
    # Row 2 is the first SKU. T = EXT QTY, Q = EXT Cost.
    assert sheet["T2"].fill.start_color.rgb.endswith(MISMATCH_FILL.start_color.rgb)
    assert sheet["T2"].comment.text == "Expected 136 (CTNS x CSPK)"
    assert sheet["Q2"].comment.text == "Expected 6767.800 (Cost x EXT QTY)"
    assert sheet["T3"].comment is None  # the second SKU is fine


def test_missing_upc_raises():
    broken = SAMPLE.replace("    0091120032129\n", "\n")
    try:
        parse_items(broken)
    except ValueError as error:
        assert "71097188" in str(error)
    else:
        raise AssertionError("expected ValueError")


def test_output_matches_template(tmp_path):
    items = parse_items(SAMPLE)
    rows = build_rows("purchase_order_sample.txt", parse_header(SAMPLE), items)
    out_path = tmp_path / "out.xlsx"
    write_xlsx(rows, [validate_item(item) for item in items], out_path)

    template_headers = [cell.value for cell in load_workbook(TEMPLATE_PATH).active[1]]
    workbook = load_workbook(out_path)
    sheet = workbook.active
    output_headers = [cell.value for cell in sheet[1]]

    assert output_headers == template_headers  # exactly the template's columns, nothing added
    assert workbook.sheetnames == ["Sheet1"]
    assert sheet.max_row == 13  # header + 12 SKUs
    assert sheet["L2"].value == "0091120032129"  # UPC stored as text
